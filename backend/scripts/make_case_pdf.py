# -*- coding: utf-8 -*-
"""生成「实操案例：用 Coding Agent 交付本项目」的提交用 PDF。

管线：agent_case.json → python-docx 组稿 → Word COM 导出 PDF → pypdf 回验文本层。

回验是硬性的一步：上一次提交被退回，原因是附件提取出来只有网页脚本代码、
读不到正文。本脚本导出后会把 PDF 的文本层重新抽出来自查——
中文字数、是否含预期标题、脚本特征占比，任一不合格就直接报错。

前置：先跑 scripts/export_agent_case.py 生成数据，
      docs/competition/shots/ 内有界面截图。
用法（工作目录必须是 backend）：
    .\\venv\\Scripts\\python.exe -u scripts\\make_case_pdf.py
"""
from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPO = Path(__file__).resolve().parents[2]
BACKEND = REPO / "backend"
DOCS = REPO / "docs" / "competition"
SHOTS = DOCS / "shots"

sys.path.insert(0, str(Path(__file__).resolve().parent))
from export_agent_case import SECRET_PATTERNS, redact  # noqa: E402  脱敏口径与导出器共用同一份规则

DATA = BACKEND / "data" / "eval" / "agent_case.json"
# 验证输出优先取仓库内的 UTF-8 存档（可复现），退回到最近一次运行的临时日志
VERIFY_LOG_REPO = DOCS / "attachments" / "verify_all_output.txt"
VERIFY_LOG_TMP = Path(r"D:\workspace\12345工单热线\verify_all_last.txt")


def read_text_any(path: Path) -> str:
    """按 BOM/内容猜编码读取。

    PowerShell 的 Tee-Object 在不同宿主下会写成 UTF-16LE，
    按 UTF-8 硬读会得到乱码并让后续关键词匹配全部失效。
    """
    raw = path.read_bytes()
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16", "replace")
    if raw[:3] == b"\xef\xbb\xbf":
        return raw.decode("utf-8-sig", "replace")
    for enc in ("utf-8", "gbk"):
        try:
            return raw.decode(enc)
        except Exception:
            continue
    return raw.decode("utf-8", "replace")

OUT_DOCX = DOCS / "11_实操案例_CodingAgent使用过程.docx"
OUT_PDF = DOCS / "11_实操案例_CodingAgent使用过程.pdf"

FONT = "微软雅黑"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)      # 深蓝：标题
INK = RGBColor(0x22, 0x22, 0x22)         # 正文近黑
MUTED = RGBColor(0x66, 0x66, 0x66)

SCRIPT_SMELL = ("<script", "</script", "function(", "document.getElementById",
                "window.__", "var ", "=>{", "console.log", "addEventListener")


# ---------------------------------------------------------------- 排版工具
def style_font(style, size=None, bold=None, color=None) -> None:
    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), FONT)
    rf.set(qn("w:ascii"), FONT)
    rf.set(qn("w:hAnsi"), FONT)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def setup(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)

    style_font(doc.styles["Normal"], size=10.5, color=INK)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.32

    style_font(doc.styles["Heading 1"], size=16, bold=True, color=ACCENT)
    style_font(doc.styles["Heading 2"], size=13, bold=True, color=ACCENT)
    style_font(doc.styles["Heading 3"], size=11.5, bold=True, color=RGBColor(0x33, 0x44, 0x55))
    for h, before in (("Heading 1", 16), ("Heading 2", 12), ("Heading 3", 10)):
        pf = doc.styles[h].paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(6)
        pf.keep_with_next = True


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    return p


def h2(doc, text):
    return doc.add_heading(text, level=2)


def h3(doc, text):
    return doc.add_heading(text, level=3)


def para(doc, text, *, size=10.5, color=None, bold=False, italic=False,
         align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(redact(str(text)))
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.color.rgb = color or INK
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(doc, text, *, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(redact(str(text)))
    r.font.size = Pt(size)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    p.paragraph_format.space_after = Pt(3)
    return p


def shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_widths(table, widths_cm) -> None:
    """强制固定列宽：Word 默认 autofit 会忽略设定值，需要显式 tblLayout=fixed。"""
    table.autofit = False
    tblpr = table._tbl.tblPr
    for tag in ("w:tblLayout",):
        for el in tblpr.findall(qn(tag)):
            tblpr.remove(el)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
            if i < len(widths_cm):
                gc.set(qn("w:w"), str(int(widths_cm[i] * 567)))
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)


def make_table(doc, header, rows, widths, *, size=9, head_fill="E8EEF4",
               max_rows=None):
    rows = rows[:max_rows] if max_rows else rows
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htext in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(str(htext))
        r.font.bold = True
        r.font.size = Pt(size)
        r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        r.font.color.rgb = ACCENT
        p.paragraph_format.space_after = Pt(0)
        shade(cell, head_fill)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            if ci >= len(header):
                break
            cell = t.rows[ri].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(redact("" if val is None else str(val)))
            r.font.size = Pt(size)
            r.font.name = FONT
            r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            p.paragraph_format.space_after = Pt(0)
    set_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def picture(doc, path: Path, caption: str, width_cm=16.4, max_height_cm=19.0) -> None:
    """插图。高图按高度上限反算宽度，避免细长截图撑爆版面。"""
    if not path.exists():
        return
    try:
        from PIL import Image
        with Image.open(path) as im:
            w_px, h_px = im.size
        if h_px / w_px * width_cm > max_height_cm:
            width_cm = max_height_cm * w_px / h_px
    except Exception:
        pass
    doc.add_picture(str(path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = para(doc, caption, size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER,
             space_after=12)
    p.runs[0].font.italic = True


# ---------------------------------------------------------------- 正文组稿
def build(d: dict) -> Document:
    doc = Document()
    setup(doc)
    tot = d["totals"]
    span = d["span"]
    agent = d["agent"]

    # ---------- 封面 ----------
    para(doc, "", space_after=60)
    para(doc, "实操案例", size=30, bold=True, color=ACCENT,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    para(doc, "用 Coding Agent 从零交付「12345 热线工单智能体」",
         size=14, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
    para(doc, "任务拆解 · Prompt 设计 · 多轮调试 · 代码交付　完整过程记录",
         size=11, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=40)
    para(doc, f"Coding Agent：{agent['harness']}（模型 {agent['model']}）",
         size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    para(doc, f"时间跨度：{span['firstDay']} — {span['lastDay']}，共 {span['days']} 个工作日",
         size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    para(doc, f"规模：{tot['turns']} 轮对话 / {tot['steps']} 个执行步 / {tot['toolCalls']} 次工具调用",
         size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    para(doc, f"生成时间：{d['generatedAt']}",
         size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------- 〇、材料说明 ----------
    h1(doc, "〇、材料说明")
    para(doc, "本材料是一份可直接阅读、检索、复制的完整正文文档，不是网页另存或截图拼贴。"
              "全文由真实会话记录导出后组稿，导出环节包含「PDF 文本层回验」："
              "从生成的 PDF 反向抽取文字，确认抽到的是正文而不是网页脚本代码。")
    para(doc, "所有统计数字与报错原文均来自 Coding Agent 的本地会话记录"
              f"（{agent['harness']}，工作区 {agent['workspace']}）。"
              "会话记录是 Agent 与开发者交互时逐事件写入的原始日志"
              "（类型 + 时间戳 + 内容，zstd 压缩的 JSONL），并非事后补写的说明文档，"
              "因此可以作为「在实际项目中使用 Coding Agent」的过程证据。")
    para(doc, "复现方式：仓库内 backend/scripts/export_agent_case.py 负责从会话记录导出过程数据，"
              "make_case_pdf.py 负责由数据生成这份 PDF 并回验文本层。"
              "更换机器后，只要会话记录在，本文即可原样重新生成。")

    # ---------- 一、工作方式与规模 ----------
    h1(doc, "一、工作方式与规模")
    para(doc, "本项目的开发方式不是「让模型写一段代码」，而是把 Coding Agent 当作执行主体："
              "由开发者给出目标与约束，Agent 自行拆解任务、读写代码、执行命令、跑验证、"
              "根据报错定位并修复，循环推进到可交付状态。")
    make_table(
        doc,
        ["项目", "数值", "说明"],
        [
            ["对话轮次", tot["turns"], "开发者每次下达目标算一轮"],
            ["执行步", tot["steps"], "Agent 内部的推理-行动步"],
            ["工具调用", tot["toolCalls"], "Agent 实际执行的动作总数"],
            ["命令执行", dict(d["tools"]).get("pwsh", 0), "跑脚本、起服务、跑验证、查环境"],
            ["文件编辑", tot["edits"], "对既有文件做定点修改"],
            ["文件新建", tot["writes"], "新增脚本、模块、文档"],
            ["文件读取", dict(d["tools"]).get("read", 0), "读代码后再改，避免盲改"],
            ["内容检索", dict(d["tools"]).get("grep", 0) + dict(d["tools"]).get("glob", 0), "跨文件定位实现"],
            ["覆盖文件数", tot["filesTouched"], "后端 + 前端 + 脚本 + 文档"],
            ["开发者指令", tot["instructions"], "含中途追加与纠正的要求"],
            ["任务拆解计划", tot["plans"], "Agent 维护的待办清单版本数"],
            ["捕获工具失败", tot["errors"], "程序异常 %d、工具拒绝 %d、命令非零退出 %d、校验未过 %d"
             % (tot["errorKinds"].get("异常", 0), tot["errorKinds"].get("工具拒绝", 0),
                tot["errorKinds"].get("非零退出", 0), tot["errorKinds"].get("校验未过", 0))],
            ["模型层自动重试", tot["retries"], "网络/超时导致的自动退避重试"],
            ["时间跨度", f"{span['days']} 天", f"{span['firstDay']} — {span['lastDay']}"],
        ],
        [3.4, 2.6, 10.4],
    )
    para(doc, "取证说明：Agent 的续接会话会重放上一段历史，直接统计记录条数会虚高。"
              f"本文按「事件类型 + 时间戳 + 内容指纹」跨会话去重，剔除了 {tot['replayedDuplicatesDropped']} 条"
              "重放记录，上表对应真实发生过的动作。", size=9.5, color=MUTED)

    # ---------- 二、任务拆解 ----------
    h1(doc, "二、任务拆解：先出计划，再逐项推进")
    para(doc, f"面对「做一个 12345 工单智能体」这种颗粒度很大的目标，Agent 没有直接动手写代码，"
              f"而是先把它拆成可验证的具体动作，并在整个过程中维护了 {len(d['plans'])} 版任务清单"
              f"（todo 事件的真实快照），每完成一项就更新状态。")
    if d["plans"]:
        first = d["plans"][0]
        h3(doc, f"第 1 版计划（{first['time']}，共 {first['total']} 项）")
        for i, x in enumerate(first["todos"], 1):
            mark = {"completed": "已完成", "in_progress": "进行中", "pending": "待办"}.get(
                x["status"], x["status"])
            bullet(doc, f"{i}. {x['content']}（{mark}）")
    h3(doc, "计划的后续演进")
    para(doc, "下表每行是一次计划更新，可以看到目标从「环境搭起来」逐步推进到"
              "「三支柱能力」「合规审查」「验证与材料」：", size=10)
    rows = []
    prev: set = set()
    for p in d["plans"]:
        cur = {x["content"] for x in p["todos"]}
        new = [c for c in cur if c not in prev]
        prev = cur
        shown = "；".join(new[:2]) if new else "（状态推进）"
        if len(new) > 2:
            shown += f" 等 {len(new)} 项"
        rows.append([p["time"], p.get("turn") or "—", f"{p['done']}/{p['total']}", shown[:70]])
    make_table(doc, ["时间", "轮次", "完成/总数", "新增或变化的条目"],
               rows, [2.1, 1.4, 2.0, 10.9], max_rows=16)

    # ---------- 三、Prompt 设计 ----------
    h1(doc, "三、Prompt 设计：开发者指令如何演进")
    ins = d["instructions"]
    para(doc, f"全程共 {len(ins)} 条开发者指令。与 Agent 协作的关键不是一次把需求写全，"
              f"而是让它先把现状讲清楚、再按约束推进。指令可以分成三类：")
    bullet(doc, "目标类：给出要达成的业务结果（例如「设计三支柱，要能实际解决问题」）。")
    bullet(doc, "约束类：给出不可违背的边界（例如「不要把 key 提交」「预留运营商的对接接口」）。")
    bullet(doc, "纠正类：看到结果后给出方向修正（例如「一股 AI 味道，再优化界面」）。")
    para(doc, "以下是按时间顺序节选的原始指令（未改写措辞）：", size=10)

    def short(t, n=150):
        t = re.sub(r"\s+", " ", t or "").strip()
        return t if len(t) <= n else t[:n] + "…"

    picks = ins[:8]
    fix_words = ("不要", "不对", "改成", "注意", "必须", "重新", "优化", "修正", "去掉", "别")
    extra = [x for x in ins[8:] if any(w in x["text"] for w in fix_words)][:12]
    shown_ids = {id(x) for x in picks} | {id(x) for x in extra}
    seq = sorted(picks + extra, key=lambda x: x.get("seq") or 0)
    rows = []
    for i, x in enumerate(seq, 1):
        rows.append([x["time"], x["chars"], short(x["text"], 160)])
    make_table(doc, ["时间", "字数", "开发者指令（原文）"], rows, [2.1, 1.1, 13.2], size=8.5)
    para(doc, f"（其余 {len(ins) - len(shown_ids)} 条指令见配套的完整过程记录文档 "
              f"10_实操案例_CodingAgent过程记录.md。）", size=9, color=MUTED)

    # ---------- 四、多轮调试 ----------
    h1(doc, "四、多轮调试：报错 → 定位 → 修复")
    ek = tot["errorKinds"]
    para(doc, f"整个过程捕获工具失败 {tot['errors']} 次（程序异常 {ek.get('异常', 0)}、"
              f"工具拒绝 {ek.get('工具拒绝', 0)}、命令非零退出 {ek.get('非零退出', 0)}、"
              f"自建校验未过 {ek.get('校验未过', 0)}），另有 {tot['retries']} 次模型层自动重试。"
              f"这些失败没有被绕过或隐藏——它们是开发过程本身，也全部当轮定位并修复。")

    h3(doc, "典型故障的全文追踪")
    para(doc, "下表故障都在完整会话记录里做过全文检索，并给出首次出现的记录原文，"
              "以保证叙述与记录一致：", size=10)
    rows = []
    for k in d["keyIssues"]:
        f = k["first"]
        rows.append([k["label"], k["count"], f["time"], short(f["quote"], 105)])
    make_table(doc, ["故障", "命中", "首次出现", "记录原文（摘自会话记录）"],
               rows, [3.3, 1.0, 2.4, 9.7], size=8.5)

    if d["retries"]:
        h3(doc, "模型层自动重试")
        codes: dict[str, int] = {}
        sample: dict[str, str] = {}
        for r in d["retries"]:
            c = r["code"] or "?"
            codes[c] = codes.get(c, 0) + 1
            sample.setdefault(c, r["message"])
        rows = [[f"`{c}`".strip("`"), n, short(sample[c], 90)] for c, n in
                sorted(codes.items(), key=lambda kv: -kv[1])]
        make_table(doc, ["错误码", "次数", "典型信息"], rows, [2.6, 1.4, 12.4], size=8.5)
        para(doc, "这些是长任务里的网络与超时抖动。Agent 按退避策略自动重试并继续，"
                  "因此十几天跨度的工作没有被一次网络中断打断。", size=9.5, color=MUTED)

    h3(doc, "一次典型调试回合（原文摘录）")
    hard = [e for e in d["errors"] if e["kind"] == "异常"]
    if hard:
        e0 = hard[1] if len(hard) > 1 else hard[0]
        para(doc, f"时间 {e0['time']}（第 {e0['turn']} 轮）：Agent 执行命令得到异常，"
                  f"摘要为「{short(e0['summary'], 110)}」；"
                  f"紧接着的动作是「{short(e0['fix'], 70)}」。"
                  f"这类「跑起来 → 看报错 → 改代码/换方案 → 再跑」的回路在记录中反复出现。", size=10)

    # ---------- 五、代码交付 ----------
    h1(doc, "五、代码交付：Agent 实际改动的文件")
    para(doc, f"Agent 通过编辑/新建工具直接交付代码，累计编辑 {tot['edits']} 次、"
              f"新建 {tot['writes']} 次，覆盖 {tot['filesTouched']} 个文件。改动最集中的文件如下：")
    rows = [[f["path"], f["edit"], f["write"]] for f in d["files"][:18]]
    make_table(doc, ["文件", "编辑", "新建"], rows, [12.6, 1.9, 1.9], size=8.5)

    ext: dict[str, int] = {}
    for f in d["files"]:
        e_ = Path(f["path"]).suffix or "(无扩展名)"
        ext[e_] = ext.get(e_, 0) + f["edit"] + f["write"]
    h3(doc, "交付物的类型分布")
    rows = [[k, v] for k, v in sorted(ext.items(), key=lambda kv: -kv[1])[:10]]
    make_table(doc, ["类型", "次数"], rows, [3.0, 2.0], size=9)

    h3(doc, "落库：git 提交记录")
    para(doc, "上述改动以提交形式进入版本库，构成可追溯的交付链路：", size=10)
    rows = []
    for line in d["commits"][:22]:
        parts = line.split("|")
        if len(parts) == 3:
            rows.append([parts[0], parts[1], short(parts[2], 80)])
    make_table(doc, ["提交", "时间", "说明"], rows, [2.2, 2.2, 12.0], size=8.5)
    para(doc, f"仓库累计 {len(d['commits'])} 条提交。", size=9, color=MUTED)

    verify_log = next((p for p in (VERIFY_LOG_REPO, VERIFY_LOG_TMP) if p.exists()), None)
    if verify_log is not None:
        text = read_text_any(verify_log)
        lines = [x for x in text.splitlines() if re.search(r"(通过|失败|PASS|FAIL|✓|✗)", x)]
        if lines:
            h3(doc, "交付前的全量验证结果")
            para(doc, "提交材料前，Agent 运行了仓库内的验证套件（scripts/verify_all.py），"
                      "覆盖服务、接口、评测与界面四层，实际输出为：", size=10)
            rows = [[short(x.strip(), 96)] for x in lines[-20:]]
            make_table(doc, ["验证输出（节选末段）"], rows, [16.0], size=8.5)
            if "全部通过" in text:
                para(doc, "总体结论：全部通过。完整输出见 docs/competition/attachments/verify_all_output.txt。",
                     size=9.5, color=MUTED)
        else:
            para(doc, "（验证日志存在但未匹配到结果行，已跳过该节，避免写入未经核实的内容。）",
                 size=9, color=MUTED)

    # ---------- 六、界面实证 ----------
    h1(doc, "六、界面实证")
    para(doc, "下面四张截图由脚本 backend/scripts/capture_case_shots.py 用无头浏览器"
              "访问本机运行的系统自动截取（整页截图，含视口外内容），"
              "对应前端当前构建产物，可在本机一键复现。")
    for name, cap in (
        ("case_01_workbench_detail.png",
         "图 1　工作台：左侧案件队列（30 条），右侧案件详情。可见五节点链路、"
         "急件分级（决策模型判定）、办理时限倒计时、承办单位与属地推荐、诉求治理提示。"),
        ("case_02_flow_board.png",
         "图 2　流转看板：工单在「受理—核实—立案—派遣—办理—核查—办结—回访—归档」之间的状态分布与时限。"),
        ("case_03_knowledge.png",
         "图 3　知识库：口径层与政策依据，可从已办案件蒸馏口径并保留修订历史。"),
        ("case_04_account_audit.png",
         "图 4　账号与审计：角色权限与操作审计（轻量 RBAC，可通过配置开关渐进启用）。"),
    ):
        picture(doc, SHOTS / name, cap)

    # ---------- 七、桌面实操截图 ----------
    h1(doc, "七、桌面实操截图")
    para(doc, "下面这张是 AI 工具在桌面上的真实运行窗口，可直接看到本次项目的会话、"
              "工具调用卡片、所用模型与权限模式。截图由 scripts/capture_desktop_shots.ps1 取得："
              "用 Win32 PrintWindow 让目标窗口渲染自身，不切换前台焦点，"
              "因此不会打断正在进行的会话，也不会把桌面上其它窗口拍进来。")
    picture(doc, SHOTS / "desktop_01_DeepSeekHarness.png",
            "图 5　桌面上的 Coding Agent 工具窗口：左侧为工作区列表（本项目 12345工单热线），"
            "中间为该项目的会话与工具调用记录，底部为模型、权限模式与运行计数。",
            max_height_cm=17.5)
    para(doc, "取证取舍说明：脚本按窗口标题匹配抓图，本次桌面上另有 PowerShell 与 VS Code 窗口，"
              "但它们打开的是其它项目（与本提交无关），为避免混入无关信息、也避免泄露其它项目内容，"
              "这两张已排除，只保留与本项目直接相关的窗口。窗口内图像不做文字提取，"
              "因此本节的密钥检查以人工目视确认为准（已确认无密钥）。",
         size=9.5, color=MUTED)

    # ---------- 八、可复核性 ----------
    h1(doc, "八、可复核性与边界")
    para(doc, "为便于核验，本材料的每个数字都对应仓库内的脚本与数据文件：", size=10)
    make_table(
        doc,
        ["内容", "来源"],
        [
            ["会话过程数据", "backend/scripts/export_agent_case.py → backend/data/eval/agent_case.json"],
            ["本文 PDF", "backend/scripts/make_case_pdf.py（docx 组稿 → Word 导出 → 文本层回验）"],
            ["完整过程记录", "docs/competition/10_实操案例_CodingAgent过程记录.md"],
            ["界面截图", "backend/scripts/capture_case_shots.py → docs/competition/shots/"],
            ["桌面截图", "backend/scripts/capture_desktop_shots.ps1 → docs/competition/shots/desktop_*.png"],
            ["全量验证", "backend/scripts/verify_all.py"],
        ],
        [3.4, 12.9],
    )
    h3(doc, "需要如实说明的边界")
    bullet(doc, "会话记录产生于开发者本机，与仓库内的代码、提交、验证输出三者可交叉印证，"
                "但它本身不能证明「没有人工代写」；本材料的作用是提供可核查的过程事实。")
    bullet(doc, "「工具失败次数」包含命令非零退出与 Agent 自身的编辑定位失配，"
                "不等于产品缺陷数；分类统计已在第一章按类别拆分。")
    bullet(doc, "指标评测样本为自建集合（分类 18 例、隐患 12 例、合规 20 例），"
                "不是公开留出集，相关阈值在样本内拟合，数字偏乐观；"
                "这一点在 09_决策模型评测 中有专门说明。")
    bullet(doc, "桌面截图由本机窗口抓取得到，证明的是「该工具在本机桌面上被实际使用」，"
                "属于过程性证据，不等同于身份认证；如需更强的归属证据，"
                "可补充会话记录的原始文件与哈希。")

    return doc


# ---------------------------------------------------------------- 导出与回验
def docx_to_pdf(docx: Path, pdf: Path) -> tuple[bool, str]:
    if pdf.exists():
        pdf.unlink()
    script = f'''
$ErrorActionPreference = "Stop"
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {{
    $doc = $word.Documents.Open("{docx}", $false, $true)
    $doc.ExportAsFixedFormat("{pdf}", 17, $false, 0, 0, 0, 0, 0, $true, $true, 1, $true, $true, $false)
    $doc.Close(0)
    Write-Output "OK"
}} finally {{
    $word.Quit()
}}
'''
    r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                       capture_output=True, text=True, encoding="utf-8", errors="replace",
                       timeout=240)
    ok = pdf.exists() and pdf.stat().st_size > 20000
    return ok, (r.stdout or "") + (r.stderr or "")


def verify_text_layer(pdf: Path) -> tuple[bool, list[str], str]:
    from pypdf import PdfReader

    reader = PdfReader(str(pdf))
    pages = [p.extract_text() or "" for p in reader.pages]
    text = "\n".join(pages)
    cjk = len(re.findall(r"[\u4e00-\u9fff]", text))
    smell = sum(text.count(s) for s in SCRIPT_SMELL)
    leaks = []
    for pat, _ in SECRET_PATTERNS:
        for m in pat.finditer(text):
            leaks.append(m.group(0)[:24])
    report = [
        f"页数：{len(pages)}",
        f"提取字符总数：{len(text)}",
        f"中文字数：{cjk}",
        f"脚本特征片段出现次数：{smell}",
        f"疑似密钥残留：{len(leaks)} 处" + (f"（{leaks[:3]}）" if leaks else ""),
        f"是否含预期标题「实操案例」：{'是' if '实操案例' in text else '否'}",
        f"是否含预期标题「多轮调试」：{'是' if '多轮调试' in text else '否'}",
        f"是否含报错原文「database is locked」：{'是' if 'database is locked' in text else '否'}",
    ]
    ok = (cjk >= 3000 and smell <= 20 and not leaks
          and "实操案例" in text and "多轮调试" in text)
    return ok, report, text


def main() -> int:
    if not DATA.exists():
        print("缺少数据文件，请先运行 scripts/export_agent_case.py：", DATA)
        return 1
    d = json.loads(DATA.read_text(encoding="utf-8"))

    doc = build(d)
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_DOCX))
    print(f"已生成 DOCX：{OUT_DOCX}（{OUT_DOCX.stat().st_size // 1024} KB）")

    ok, log = docx_to_pdf(OUT_DOCX, OUT_PDF)
    if not ok:
        print("PDF 导出失败，转换器输出：")
        print(log[-1500:])
        return 1
    print(f"已生成 PDF：{OUT_PDF}（{OUT_PDF.stat().st_size // 1024} KB）")

    good, report, text = verify_text_layer(OUT_PDF)
    print("\n=== PDF 文本层回验 ===")
    for line in report:
        print("  " + line)
    print("\n提取正文开头 400 字（证明抽到的是正文而非脚本）：")
    print("  " + re.sub(r"\s+", " ", text)[:400])

    if not good:
        print("\n文本层回验未通过：生成物可能仍无法被评审系统读出正文。")
        return 1
    print("\n文本层回验通过：PDF 可直接检索复制中文正文。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
